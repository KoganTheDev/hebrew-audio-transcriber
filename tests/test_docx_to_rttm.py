"""
Tests for tests/eval/docx_to_rttm.py - the colour-coded .docx -> RTTM aligner.

All fixtures here are small documents built in-memory with python-docx, not
the real test#2/test#3 files (those need real audio and are exercised by
hand, see docs/TESTING.md). Each test targets one piece the module docstring
calls out as load-bearing: colour-to-speaker mapping from the header, a
leading timestamp stripped before reading a fused line's colour, majority-
colour fallback when a line's runs disagree, block parsing with the bidi
isolate characters, and the SequenceMatcher alignment/windowing logic.
"""

from dataclasses import dataclass

import docx
from docx.shared import RGBColor

from tests.eval.docx_to_rttm import (
    Block,
    align_blocks,
    parse_docx,
    parse_speaker_header,
)

PURPLE = RGBColor(0xA0, 0x2B, 0x93)


def _add_run(paragraph, text, color=None):
    run = paragraph.add_run(text)
    if color is not None:
        run.font.color.rgb = color
    return run


def _header_paragraph(document, first_name="אבי", second_name="נאור"):
    p = document.add_paragraph()
    _add_run(p, "הדוברים:")
    _add_run(p, " ")
    _add_run(p, first_name + " ", color=PURPLE)
    _add_run(p, second_name)
    return p


def _timestamp_paragraph(document, start_label, end_label):
    p = document.add_paragraph()
    _add_run(p, "⁦")
    _add_run(p, f"{start_label} - {end_label}")
    _add_run(p, "⁩")
    _add_run(p, " ")
    return p


def _speech_paragraph(document, text, color=None):
    p = document.add_paragraph()
    _add_run(p, text, color=color)
    return p


class TestParseSpeakerHeader:
    def test_maps_purple_to_first_named_speaker_and_default_to_second(self):
        document = docx.Document()
        header = _header_paragraph(document)
        mapping = parse_speaker_header(header)
        assert mapping == {"A02B93": "אבי", "DEFAULT": "נאור"}


class TestParseDocxHeaderAssertion:
    def test_accepts_the_expected_layout(self, tmp_path):
        document = docx.Document()
        _header_paragraph(document)
        _timestamp_paragraph(document, "0:00", "0:05")
        _speech_paragraph(document, "שלום", color=PURPLE)
        path = tmp_path / "ok.docx"
        document.save(str(path))

        colors, blocks = parse_docx(str(path))
        assert colors == {"A02B93": "אבי", "DEFAULT": "נאור"}
        assert len(blocks) == 1
        assert blocks[0].lines == [("אבי", "שלום")]

    def test_rejects_a_header_that_does_not_match_the_expected_colours(self, tmp_path):
        document = docx.Document()
        # Second speaker is not נאור - the expectation this app's two real
        # source files both satisfy, so a differently-authored docx must
        # fail loudly here instead of silently mislabelling every line.
        _header_paragraph(document, second_name="דנה")
        _timestamp_paragraph(document, "0:00", "0:05")
        _speech_paragraph(document, "שלום", color=PURPLE)
        path = tmp_path / "bad.docx"
        document.save(str(path))

        try:
            parse_docx(str(path))
            raised = False
        except ValueError:
            raised = True
        assert raised


class TestBlockParsing:
    def test_bidi_isolate_timestamp_starts_a_new_block(self, tmp_path):
        document = docx.Document()
        _header_paragraph(document)
        _timestamp_paragraph(document, "0:00", "0:24")
        _speech_paragraph(document, "שורה ראשונה", color=PURPLE)
        _timestamp_paragraph(document, "0:24", "0:48")
        _speech_paragraph(document, "שורה שנייה")
        path = tmp_path / "blocks.docx"
        document.save(str(path))

        _colors, blocks = parse_docx(str(path))
        assert [(b.start, b.end) for b in blocks] == [(0.0, 24.0), (24.0, 48.0)]
        assert blocks[0].lines == [("אבי", "שורה ראשונה")]
        assert blocks[1].lines == [("נאור", "שורה שנייה")]

    def test_blank_paragraphs_between_lines_are_skipped(self, tmp_path):
        document = docx.Document()
        _header_paragraph(document)
        _timestamp_paragraph(document, "0:00", "0:24")
        _speech_paragraph(document, "אחת", color=PURPLE)
        document.add_paragraph("")
        _speech_paragraph(document, "שתיים", color=PURPLE)
        path = tmp_path / "blank.docx"
        document.save(str(path))

        _colors, blocks = parse_docx(str(path))
        assert len(blocks) == 1
        assert blocks[0].lines == [("אבי", "אחת"), ("אבי", "שתיים")]

    def test_leading_timestamp_fused_onto_a_speech_line_is_stripped(self, tmp_path):
        """test#3's paragraph 79: a block marker with no blank paragraph
        after it, followed immediately (same paragraph) by the speech text -
        the marker's own runs must not pollute the speech line's colour."""
        document = docx.Document()
        _header_paragraph(document)
        p = document.add_paragraph()
        _add_run(p, "⁦")
        _add_run(p, "4:11 - 4:30")
        _add_run(p, "⁩")
        _add_run(p, " ")
        _add_run(p, "\n")
        _add_run(p, "הוא אמר משהו", color=PURPLE)
        path = tmp_path / "fused.docx"
        document.save(str(path))

        _colors, blocks = parse_docx(str(path))
        assert len(blocks) == 1
        assert blocks[0].start == 4 * 60 + 11
        assert blocks[0].end == 4 * 60 + 30
        assert blocks[0].lines == [("אבי", "הוא אמר משהו")]

    def test_majority_colour_fallback_on_mixed_runs(self, tmp_path):
        """test#3's paragraph 218: one trailing punctuation run keeps the
        paragraph's default colour while every other run is coloured for
        the real speaker - the line as a whole must follow the majority."""
        document = docx.Document()
        _header_paragraph(document)
        _timestamp_paragraph(document, "0:00", "0:24")
        p = document.add_paragraph()
        _add_run(p, "משפט ארוך למדי כאן", color=PURPLE)
        _add_run(p, ".")  # default colour, much shorter than the rest
        path = tmp_path / "mixed.docx"
        document.save(str(path))

        _colors, blocks = parse_docx(str(path))
        assert blocks[0].lines == [("אבי", "משפט ארוך למדי כאן.")]


@dataclass
class _Word:
    start: float
    end: float
    text: str


@dataclass
class _Segment:
    start: float
    end: float
    text: str
    words: list


def _segment(words):
    return _Segment(
        start=words[0].start if words else 0.0,
        end=words[-1].end if words else 0.0,
        text=" ".join(w.text for w in words),
        words=words,
    )


class TestAlignBlocks:
    def test_matches_a_line_to_its_windowed_words(self):
        block = Block(start=0.0, end=10.0, lines=[("אבי", "שלום עולם")])
        words = [
            _Word(0.5, 1.0, "שלום"),
            _Word(1.1, 1.6, "עולם"),
        ]
        segments = [_segment(words)]

        bubbles, matched, dropped = align_blocks([block], segments, pad=3.0)
        assert matched == 1
        assert dropped == 0
        assert len(bubbles) == 1
        assert bubbles[0].speaker == "אבי"
        assert bubbles[0].start == 0.5
        assert bubbles[0].end == 1.6

    def test_words_outside_the_padded_window_are_ignored(self):
        """A word that matches the text but sits well outside the block's
        window (even padded) must not be used - the window is what makes
        matching tractable and it is a hard constraint, not a hint."""
        block = Block(start=100.0, end=110.0, lines=[("אבי", "שלום")])
        words = [_Word(0.0, 0.5, "שלום")]  # same word, far outside the window
        segments = [_segment(words)]

        bubbles, matched, dropped = align_blocks([block], segments, pad=3.0)
        assert matched == 0
        assert dropped == 1
        assert bubbles == []

    def test_low_coverage_line_is_dropped_not_guessed(self):
        block = Block(
            start=0.0,
            end=10.0,
            lines=[("אבי", "מילים שלא הוקלטו בכלל ולא יימצאו")],
        )
        words = [_Word(0.5, 1.0, "משהו"), _Word(1.1, 1.6, "אחר")]
        segments = [_segment(words)]

        bubbles, matched, dropped = align_blocks([block], segments, pad=3.0, min_line_coverage=0.5)
        assert matched == 0
        assert dropped == 1
        assert bubbles == []

    def test_two_lines_in_one_block_get_distinct_spans(self):
        block = Block(
            start=0.0,
            end=20.0,
            lines=[("אבי", "מה שלומך"), ("נאור", "טוב תודה")],
        )
        words = [
            _Word(1.0, 1.4, "מה"),
            _Word(1.5, 2.0, "שלומך"),
            _Word(5.0, 5.3, "טוב"),
            _Word(5.4, 5.9, "תודה"),
        ]
        segments = [_segment(words)]

        bubbles, matched, dropped = align_blocks([block], segments, pad=3.0)
        assert matched == 2 and dropped == 0
        by_speaker = {b.speaker: (b.start, b.end) for b in bubbles}
        assert by_speaker["אבי"] == (1.0, 2.0)
        assert by_speaker["נאור"] == (5.0, 5.9)
